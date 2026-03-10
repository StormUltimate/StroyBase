#!/usr/bin/env python3
"""
Импорт работ и ежедневного выполнения из Excel в таблицы works и work_progress.

Поддержка объединённых ячеек (openpyxl): колонки A:B, объединённые как одна (Наименование),
C=Объём, D=ед. изм., E=Процент, F=План. дата, G=Выполнение (сумма), H+ = даты по дням.

Использование:
  python scripts/import_works_from_excel.py work1.xlsx --project-id 9 [--dry-run]
  python scripts/import_works_from_excel.py path/to/work1.xlsx --project-id 9 --building-sheet "6 корпус:Корпус 6"
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime, date, timedelta
import difflib

EXCEL_ORIGIN = date(1899, 12, 30)


def excel_serial_to_date(serial) -> date | None:
    if serial is None:
        return None
    try:
        if isinstance(serial, (int, float)):
            return EXCEL_ORIGIN + timedelta(days=int(serial))
        if isinstance(serial, datetime):
            return serial.date()
        if isinstance(serial, date):
            return serial
        return EXCEL_ORIGIN + timedelta(days=int(float(serial)))
    except (ValueError, TypeError):
        return None


def is_excel_date_column(value) -> bool:
    try:
        return 40000 <= float(value) <= 50000
    except (ValueError, TypeError):
        return False


def normalize_name(value: str) -> str:
    """Нормализовать наименование работы для поиска/сравнения.

    - lower()
    - обрезать пробелы
    - схлопнуть повторяющиеся пробелы
    - убрать лишнюю пунктуацию по краям
    """
    if not value:
        return ""
    s = str(value).strip().lower()
    # Удаляем типичную пунктуацию по краям
    s = s.strip(".,;:!\"'«»()[]{}")
    # Схлопываем последовательности пробелов
    s = " ".join(s.split())
    return s


def _build_merge_map(ws):
    """По листу openpyxl строим словарь (row_1based, col_1based) -> (min_row, min_col) для верхней левой ячейки объединения."""
    merge_map = {}
    for merged_range in ws.merged_cells.ranges:
        min_r = merged_range.min_row
        min_c = merged_range.min_col
        for r in range(merged_range.min_row, merged_range.max_row + 1):
            for c in range(merged_range.min_col, merged_range.max_col + 1):
                merge_map[(r, c)] = (min_r, min_c)
    return merge_map


def _effective_rows(ws, max_rows=None):
    """Читаем строки листа с учётом объединённых ячеек: значение берётся из верхней левой ячейки объединения."""
    merge_map = _build_merge_map(ws)
    rows = list(ws.iter_rows(values_only=True))
    if max_rows:
        rows = rows[:max_rows]
    result = []
    for ri, row in enumerate(rows):
        if not row:
            result.append(list(row))
            continue
        effective = []
        for ci in range(len(row)):
            r1, c1 = ri + 1, ci + 1
            if (r1, c1) in merge_map:
                mr, mc = merge_map[(r1, c1)]
                # 0-based индексы
                val = rows[mr - 1][mc - 1] if mr - 1 < len(rows) and mc - 1 < len(rows[mr - 1]) else row[ci]
            else:
                val = row[ci]
            effective.append(val)
        result.append(effective)
    return result


def _load_row_rules(base_dir: str) -> dict[int, dict]:
    """Загрузить правила по строкам из docs/EXCEL_WORKS_TEMPLATE.json.

    Формат файла:
    {
      "row_rules": [
        {"row": 33, "type": "category", "category": "Демонтажные работы"},
        {"row": 45, "type": "category", "category": "Общестроительные работы"},
        {"row": 55, "type": "category", "category": "Инженерные системы"},
        {"row": 60, "type": "subsection", "category": "Инженерные системы", "subsection": "Устройство системы отопления"},
        {"row": 90, "type": "skip"}
      ]
    }
    """
    template_path = os.path.join(base_dir, "docs", "EXCEL_WORKS_TEMPLATE.json")
    if not os.path.isfile(template_path):
        return {}
    try:
        with open(template_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception as exc:  # noqa: BLE001
        print(f"Не удалось прочитать шаблон строк из {template_path}: {exc}", file=sys.stderr)
        return {}

    rules = {}
    for item in data.get("row_rules", []):
        try:
            row = int(item.get("row"))
        except (TypeError, ValueError):
            continue
        if row <= 0:
            continue
        rtype = (item.get("type") or "category").strip().lower()
        if rtype not in ("category", "subsection", "skip"):
            continue
        rules[row] = {
            "type": rtype,
            "category": item.get("category"),
            "subsection": item.get("subsection"),
        }
    return rules


def run_import(
    filepath: str,
    project_id: int,
    dry_run: bool = False,
    building_sheet_map: list[str] | None = None,
    clean: bool = False,
    force_overwrite: bool = False,
    verbose: bool = False,
) -> None:
    try:
        import openpyxl
    except ImportError:
        print("Установите openpyxl: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    sys.path.insert(0, base_dir)
    from app import create_app
    from app.extensions import db
    from app.models import Project, Building, Work, WorkProgress
    from sqlalchemy import text

    app = create_app()
    with app.app_context():
        project = db.session.get(Project, project_id)
        if not project:
            print(f"Проект с id={project_id} не найден.", file=sys.stderr)
            sys.exit(1)

        if clean and not dry_run:
            db.session.execute(text("DELETE FROM work_progress WHERE work_id IN (SELECT id FROM works WHERE project_id = :pid)"), {"pid": project_id})
            db.session.execute(text("DELETE FROM works WHERE project_id = :pid"), {"pid": project_id})
            db.session.commit()
            print("Таблицы works и work_progress очищены по проекту.")

        buildings = {b.name.strip().lower(): b for b in Building.query.filter_by(project_id=project_id).all()}
        if not buildings:
            print("У проекта нет корпусов. Создайте корпуса в приложении.", file=sys.stderr)
            sys.exit(1)

        row_rules = _load_row_rules(base_dir)

        wb = openpyxl.load_workbook(filepath, read_only=False, data_only=True)
        sheet_to_building = {}
        if building_sheet_map:
            for pair in building_sheet_map:
                if ":" in pair:
                    sh, bd = pair.split(":", 1)
                    sheet_to_building[sh.strip()] = bd.strip()

        created_works = 0
        created_progress = 0

        # Кеш существующих работ по корпусам для нормализованного поиска и поиска похожих имён
        all_existing_works = (
            Work.query.filter(Work.project_id == project_id)
            .order_by(Work.building_id, Work.id)
            .all()
        )
        by_building_and_normname: dict[tuple[int | None, str], Work] = {}
        by_building_names: dict[int | None, list[tuple[str, Work]]] = {}
        for w in all_existing_works:
            norm = normalize_name(w.name or "")
            key = (w.building_id, norm)
            if norm:
                by_building_and_normname[key] = w
                by_building_names.setdefault(w.building_id, []).append((norm, w))

        for sheet_name in wb.sheetnames:
            if sheet_name == "Общая":
                continue
            sheet_lower = sheet_name.strip().lower()
            if "корпус" not in sheet_lower and sheet_name not in sheet_to_building:
                continue

            building_name = sheet_to_building.get(sheet_name) or sheet_name.strip()
            building = buildings.get(building_name.lower())
            if not building:
                for bkey, b in buildings.items():
                    if bkey in sheet_lower or sheet_lower in bkey:
                        building = b
                        break
            if not building:
                digits = re.findall(r"\d+", sheet_name)
                if digits:
                    num = digits[0]
                    for bkey, b in buildings.items():
                        if num in bkey or (b.name and num in str(b.name)):
                            building = b
                            break
            if not building:
                print(f"  Пропуск листа '{sheet_name}': корпус не найден в БД.")
                continue

            ws = wb[sheet_name]
            rows = _effective_rows(ws, max_rows=500)
            if len(rows) < 31:
                continue

            header_row_idx = 29
            for ri in range(min(35, len(rows))):
                r = rows[ri]
                if not r:
                    continue
                for c in (r[:10] if len(r) >= 10 else r):
                    if c is not None and "тип работ" in str(c).lower():
                        header_row_idx = ri
                        break
                else:
                    continue
                break

            headers = rows[header_row_idx] if header_row_idx < len(rows) else []
            # A:B merged = col0 (Name), C=Volume, D=Unit, E=Percent, F=Planned, G=Execution, H+= daily
            name_col = 0
            volume_col = 2
            unit_col, percent_col, planned_col, execution_col = 3, 4, 5, 6
            # Попробуем найти колонку "Категория"/"Раздел"/"Тип"/"Этап"/"Группа"
            category_col = None
            category_header_markers = ('категор', 'раздел', 'тип', 'этап', 'группа')
            for idx, h in enumerate(headers):
                if h is None:
                    continue
                h_str = str(h).strip().lower()
                if any(mark in h_str for mark in category_header_markers):
                    category_col = idx
                    break
            date_columns = []
            for col_idx in range(7, len(headers)):
                if col_idx >= len(headers):
                    break
                h = headers[col_idx]
                if isinstance(h, datetime):
                    date_columns.append((col_idx, h.date()))
                elif isinstance(h, date):
                    date_columns.append((col_idx, h))
                elif is_excel_date_column(h):
                    date_columns.append((col_idx, excel_serial_to_date(h)))

            current_category = None      # Активный раздел работ по заголовкам "Демонтажные работы" / "Общестроительные работы" / "Инженерные системы"
            current_subsection = None    # Подраздел инженерных систем (электрика, вентиляция и т.п.) по строкам без объёма внутри раздела "Инженерные системы"

            for row_idx in range(header_row_idx + 1, len(rows)):
                row = rows[row_idx]
                excel_row = row_idx + 1  # 1-based индекс строки, как в Excel

                # Правила из шаблона: жёстко задаём, какие строки являются заголовками/подразделами/служебными
                rule = row_rules.get(excel_row) if row_rules else None
                if rule:
                    rtype = rule.get("type")
                    if rtype == "category":
                        # Строка-заголовок раздела: обновляем текущий раздел, в БД ничего не пишем
                        if rule.get("category"):
                            current_category = rule["category"]
                        current_subsection = None
                        if verbose:
                            print(f"  [row {excel_row}] заголовок раздела по шаблону: category={current_category!r}")
                        continue
                    if rtype == "subsection":
                        # Строка-заголовок подраздела (обычно для инженерных систем)
                        if rule.get("category"):
                            current_category = rule["category"]
                        current_subsection = rule.get("subsection") or current_subsection
                        if verbose:
                            print(
                                f"  [row {excel_row}] заголовок подраздела по шаблону: "
                                f"category={current_category!r}, subsection={current_subsection!r}"
                            )
                        continue
                    if rtype == "skip":
                        # Строку полностью игнорируем, даже если там есть объём
                        if verbose:
                            print(f"  [row {excel_row}] строка помечена как skip в шаблоне — пропуск")
                        continue

                name = row[name_col] if name_col < len(row) else None
                if not name or (isinstance(name, str) and not name.strip()):
                    continue
                name = str(name).strip() if name else None
                if name in ("Тип работ", "Наименование системы", "ед. изм."):
                    continue

                volume = row[volume_col] if volume_col < len(row) else None
                try:
                    volume = float(volume) if volume is not None else None
                except (ValueError, TypeError):
                    volume = None
                # Если объёма нет или он ≤ 0 — это либо заголовок раздела, либо подзаголовок инженерных систем, либо техническая строка. В БД не пишем.
                if volume is None or (isinstance(volume, (int, float)) and volume <= 0):
                    name_lower = name.lower()
                    # Заголовки разделов: "Демонтажные работы" / "Общестроительные работы" / "Инженерные системы"
                    if 'демонтажные работы' in name_lower:
                        current_category = 'Демонтажные работы'
                        current_subsection = None
                    elif 'общестроительные работы' in name_lower:
                        current_category = 'Общестроительные работы'
                        current_subsection = None
                    elif 'инженерные системы' in name_lower:
                        current_category = 'Инженерные системы'
                        current_subsection = None
                    # Подзаголовки инженерных систем: строки без объёма внутри раздела "Инженерные системы"
                    elif current_category == 'Инженерные системы':
                        # Пропускаем строку с заголовками столбцов, все остальные без объёма считаем названием подраздела
                        if name_lower not in ('наименование', 'наименование системы', 'ед. изм.'):
                            current_subsection = name
                    # Все прочие строки без объёма (итоги и пр.) просто игнорируем.
                    if verbose:
                        print(f"  [row {row_idx + 1}] '{name}' без объёма — пропуск (current_category={current_category!r}, current_subsection={current_subsection!r})")
                    continue
                unit = row[unit_col] if unit_col < len(row) else None
                unit = str(unit).strip() if unit else None
                planned_val = row[planned_col] if planned_col < len(row) else None
                planned_date = excel_serial_to_date(planned_val) if planned_val is not None else None
                pct = row[percent_col] if percent_col < len(row) else None
                try:
                    percent_complete = float(pct) if pct is not None else None
                    if percent_complete is not None and 0 < percent_complete <= 1:
                        percent_complete = percent_complete * 100
                except (ValueError, TypeError):
                    percent_complete = None
                execution_total = None
                if execution_col < len(row) and row[execution_col] is not None:
                    try:
                        execution_total = float(row[execution_col])
                        if execution_total < 0:
                            execution_total = None
                    except (ValueError, TypeError):
                        pass

                has_daily = False
                for col_idx, day_date in date_columns:
                    if col_idx < len(row) and row[col_idx] is not None:
                        try:
                            if float(row[col_idx]) != 0:
                                has_daily = True
                                break
                        except (ValueError, TypeError):
                            pass
                initial_executed = None
                if not has_daily:
                    if execution_total is not None:
                        initial_executed = execution_total
                    elif percent_complete is not None and volume and volume > 0:
                        initial_executed = round(volume * (percent_complete / 100.0), 2)

                # Категория / раздел работ
                category = None
                raw_category_value = None

                # В "шаблонном" режиме (когда есть EXCEL_WORKS_TEMPLATE.json) считаем,
                # что главная истина — это активный заголовок раздела (current_category),
                # а не текст в отдельной колонке "Категория". Так строки 31–38 остаются
                # в своём разделе ("Демонтажные работы"), даже если в колонке написано
                # что‑то про инженерные системы.
                if row_rules and current_category is not None:
                    category = current_category
                else:
                    if category_col is not None and category_col < len(row):
                        raw_category_value = row[category_col]

                    if raw_category_value is not None:
                        s = str(raw_category_value).strip().lower()
                        if 'демонтаж' in s:
                            category = 'Демонтажные работы'
                        elif 'инженер' in s or 'систем' in s or 'сет' in s:
                            category = 'Инженерные системы'
                        elif 'общестро' in s or 'общие строит' in s or 'общее строит' in s:
                            category = 'Общестроительные работы'

                    # Если в явной колонке категория не задана — берём из текущего заголовка раздела
                    if category is None:
                        category = current_category or 'Общестроительные работы'

                # Подраздел инженерных систем: только для category = 'Инженерные системы'
                system_subsection = None
                if category == 'Инженерные системы':
                    # В «линейном» режиме полагаемся только на текущий подзаголовок
                    # из шаблона/заголовков, без дополнительных привязок по имени.
                    system_subsection = current_subsection

                if verbose:
                    cat_src = f"столбец #{category_col + 1}" if raw_category_value is not None and category_col is not None else "по активному заголовку раздела"
                    print(
                        f"  [row {row_idx + 1}] '{name}' — категория: '{category}', подраздел: {system_subsection!r} "
                        f"({cat_src}, значение={raw_category_value!r}, current_category={current_category!r}, current_subsection={current_subsection!r})"
                    )

                # Поиск существующей работы по проекту/корпусу/наименованию (нормализованно)
                norm_name = normalize_name(name)
                existing_work = by_building_and_normname.get((building.id, norm_name))

                similar_target_work = None
                similar_ratio = 0.0
                if not existing_work and norm_name:
                    for existing_norm, existing in by_building_names.get(building.id, []):
                        ratio = difflib.SequenceMatcher(a=norm_name, b=existing_norm).ratio()
                        if ratio > 0.85 and ratio > similar_ratio:
                            similar_ratio = ratio
                            similar_target_work = existing

                if dry_run:
                    target = existing_work or similar_target_work
                    if target:
                        has_progress = WorkProgress.query.filter_by(work_id=target.id).count() > 0
                        action = "UPDATE (по точному совпадению имени)" if existing_work else "UPDATE (по похожему имени)"
                        print(
                            f"  [dry-run] {action} Work(id={target.id}): "
                            f"volume={volume}, unit={unit}, planned={planned_date}, %={percent_complete}, "
                            f"initial_executed={initial_executed} (has_progress={has_progress})"
                        )
                    else:
                        print(
                            f"  [dry-run] INSERT Work: name={name}, volume={volume}, unit={unit}, "
                            f"planned={planned_date}, %={percent_complete}, initial_executed={initial_executed}"
                        )
                        created_works += 1
                    continue

                # Если есть только похожее имя и нет точного совпадения
                if not existing_work and similar_target_work:
                    msg = (
                        f"  [warn] Найдена похожая работа в корпусе (similarity={similar_ratio:.2f}): "
                        f"'{name}' → существующая '{similar_target_work.name}' (id={similar_target_work.id}). "
                    )
                    if not force_overwrite:
                        print(msg + "Строка из Excel пропущена (без --force-overwrite).")
                        continue
                    print(msg + "Обновляем существующую работу (--force-overwrite).")
                    existing_work = similar_target_work

                if existing_work:
                    # Работа уже есть в БД
                    has_progress = WorkProgress.query.filter_by(work_id=existing_work.id).count() > 0
                    existing_work.volume = volume
                    existing_work.unit = unit
                    existing_work.planned_completion_date = planned_date
                    existing_work.percent_complete = percent_complete
                    existing_work.category = category or existing_work.category
                    if category == 'Инженерные системы':
                        existing_work.system_subsection = system_subsection
                    # Порядок отображения в корпусе: всегда следуем порядку строк Excel
                    existing_work.sort_order = row_idx

                    if not has_daily:
                        # Обновляем initial_executed только если НЕТ записей work_progress
                        if not has_progress:
                            existing_work.initial_executed = initial_executed
                        else:
                            print(
                                f"  [warn] Работа id={existing_work.id} '{existing_work.name}' уже имеет записи "
                                f"в work_progress, initial_executed из Excel не применяется."
                            )
                    else:
                        # В Excel есть поколоночное выполнение; при наличии существующих записей по дням
                        # не добавляем новые, чтобы не задвоить объём.
                        if not has_progress:
                            for col_idx, day_date in date_columns:
                                if col_idx >= len(row):
                                    continue
                                val = row[col_idx]
                                try:
                                    daily_val = float(val) if val is not None else None
                                except (ValueError, TypeError):
                                    daily_val = None
                                if daily_val is None:
                                    continue
                                db.session.add(
                                    WorkProgress(
                                        work_id=existing_work.id,
                                        date=day_date,
                                        daily_execution=daily_val,
                                    )
                                )
                                created_progress += 1
                        else:
                            print(
                                f"  [warn] Работа id={existing_work.id} '{existing_work.name}' уже имеет записи "
                                f"в work_progress, поколоночные данные из Excel пропущены."
                            )
                else:
                    # Новая работа
                    work = Work(
                        project_id=project_id,
                        building_id=building.id,
                        name=name,
                        volume=volume,
                        unit=unit,
                        planned_completion_date=planned_date,
                        percent_complete=percent_complete,
                        initial_executed=initial_executed,
                        sort_order=row_idx,
                        category=category,
                        system_subsection=system_subsection,
                    )
                    db.session.add(work)
                    db.session.flush()
                    created_works += 1

                    for col_idx, day_date in date_columns:
                        if col_idx >= len(row):
                            continue
                        val = row[col_idx]
                        try:
                            daily_val = float(val) if val is not None else None
                        except (ValueError, TypeError):
                            daily_val = None
                        if daily_val is None:
                            continue
                        progress = WorkProgress(
                            work_id=work.id,
                            date=day_date,
                            daily_execution=daily_val,
                        )
                        db.session.add(progress)
                        created_progress += 1

        if not dry_run and (created_works > 0 or created_progress > 0):
            db.session.commit()
            print(f"Создано работ: {created_works}, записей выполнения: {created_progress}.")
        elif dry_run:
            print(f"[dry-run] Будет создано работ: {created_works}, записей выполнения: {created_progress}.")

    wb.close()


def main():
    parser = argparse.ArgumentParser(description="Импорт работ и ежедневного выполнения из Excel в works/work_progress")
    parser.add_argument("filepath", help="Путь к файлу .xlsx")
    parser.add_argument("--project-id", type=int, required=True, help="ID проекта в БД")
    parser.add_argument("--dry-run", action="store_true", help="Не записывать в БД")
    parser.add_argument("--clean", action="store_true", help="Перед импортом удалить все works и work_progress по проекту (TRUNCATE по проекту)")
    parser.add_argument("--building-sheet", action="append", dest="building_sheet_map", metavar="SHEET:BUILDING", help="Маппинг лист→корпус")
    parser.add_argument(
        "--force-overwrite",
        action="store_true",
        help="Обновлять существующие работы по похожим именам (ratio > 0.85) вместо пропуска строк",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="Подробный вывод по строкам (включая определённую категорию работ)",
    )
    args = parser.parse_args()
    if not os.path.isfile(args.filepath):
        print(f"Файл не найден: {args.filepath}", file=sys.stderr)
        sys.exit(1)
    run_import(
        args.filepath,
        args.project_id,
        dry_run=args.dry_run,
        building_sheet_map=args.building_sheet_map,
        clean=args.clean,
        force_overwrite=args.force_overwrite,
        verbose=args.verbose,
    )


if __name__ == "__main__":
    main()
