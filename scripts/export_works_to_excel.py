#!/usr/bin/env python3
"""
Экспорт работ и ежедневного выполнения из БД (works, work_progress) в Excel в формате work1.xlsx.

Структура: один лист на корпус; колонки A:B объединены под «Тип работ», C=Объём, D=ед. изм.,
E=Процент, F=Дата завершения по ГПР, G=Выполнение (сумма), H+ = даты по дням (объём за день).

Использование:
  python scripts/export_works_to_excel.py --project-id 9 -o work_export.xlsx
"""

from __future__ import annotations

import argparse
import os
import sys


def run_export(project_id: int, output_path: str):
    try:
        import openpyxl
    except ImportError:
        print("Установите openpyxl: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from app import create_app
    from app.extensions import db
    from app.models import Project, Building, Work, WorkProgress

    app = create_app()
    with app.app_context():
        project = db.session.get(Project, project_id)
        if not project:
            print(f"Проект с id={project_id} не найден.", file=sys.stderr)
            sys.exit(1)

        buildings = (
            Building.query.filter_by(project_id=project_id)
            .order_by(Building.name)
            .all()
        )
        if not buildings:
            print("У проекта нет корпусов.", file=sys.stderr)
            sys.exit(1)

        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        for building in buildings:
            works = (
                Work.query.filter_by(project_id=project_id, building_id=building.id)
                .order_by(Work.sort_order.asc().nullslast(), Work.id)
                .all()
            )

            sheet_name = building.name[:31]
            ws = wb.create_sheet(title=sheet_name)

            all_dates = set()
            work_rows = []
            for w in works:
                # Ежедневные объёмы по датам
                progress_list = w.progress.order_by(WorkProgress.date.asc()).all()
                by_date = {
                    p.date: (p.daily_execution or 0) for p in progress_list if p.date
                }
                all_dates.update(by_date.keys())

                vol = w.volume or 0
                base = float(getattr(w, "initial_executed", 0) or 0)
                daily_sum = sum(by_date.values())
                total_executed = base + daily_sum

                if vol:
                    pct = total_executed / vol * 100.0
                else:
                    pct = w.percent_complete

                work_rows.append(
                    {
                        "work": w,
                        "by_date": by_date,
                        "executed": total_executed,
                        "percent": pct,
                    }
                )
            sorted_dates = sorted(all_dates) if all_dates else []

            header_row = 30
            ws.cell(row=header_row, column=1, value="Тип работ")
            ws.merge_cells(
                start_row=header_row, start_column=1, end_row=header_row, end_column=2
            )
            ws.cell(row=header_row, column=3, value="Объем")
            ws.cell(row=header_row, column=4, value="ед. изм.")
            ws.cell(row=header_row, column=5, value="Процент")
            ws.cell(row=header_row, column=6, value="Дата завершения по ГПР")
            ws.cell(row=header_row, column=7, value="Выполнение")
            for ci, d in enumerate(sorted_dates, start=8):
                ws.cell(row=header_row, column=ci, value=d)

            for row_idx, item in enumerate(work_rows, start=header_row + 1):
                w = item["work"]
                ws.cell(row=row_idx, column=1, value=w.name or "")
                ws.merge_cells(
                    start_row=row_idx, start_column=1, end_row=row_idx, end_column=2
                )
                ws.cell(row=row_idx, column=3, value=w.volume)
                ws.cell(row=row_idx, column=4, value=w.unit or "")
                ws.cell(row=row_idx, column=5, value=item["percent"])
                ws.cell(row=row_idx, column=6, value=w.planned_completion_date)
                ws.cell(row=row_idx, column=7, value=item["executed"])
                by_date = item["by_date"]
                for ci, d in enumerate(sorted_dates, start=8):
                    ws.cell(row=row_idx, column=ci, value=by_date.get(d))

        wb.save(output_path)
        print(f"Сохранено: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Экспорт works/work_progress в Excel")
    parser.add_argument("--project-id", type=int, required=True, help="ID проекта")
    parser.add_argument("-o", "--output", required=True, help="Путь к выходному .xlsx")
    args = parser.parse_args()
    run_export(args.project_id, args.output)


if __name__ == "__main__":
    main()
